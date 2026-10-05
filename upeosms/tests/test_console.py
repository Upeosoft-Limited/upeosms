from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase

from upeosms.api import page
from upeosms.services.campaign_history import CampaignHistory
from upeosms.services.campaign_report import CampaignReport
from upeosms.services.campaign_retry import CampaignRetry
from upeosms.services.console_settings import ConsoleSettings
from upeosms.services.message_composer import MessageComposer
from upeosms.services.quick_send import QuickSend
from upeosms.services.sample_file import SampleRecipientFile
from upeosms.services.sender_profile import SenderProfile
from upeosms.utils.file_parser import _read_xlsx

TEXTSMS_KEYS = {
	"textsms_api_key": "test-key",
	"textsms_partner_id": "1",
	"textsms_sender_id": "TESTSENDER",
}
TEXTSMS_OK = {"responses": [{"response-description": "Success", "response-code": 200, "messageid": "m1"}]}


class TestMessageComposer(IntegrationTestCase):
	def test_signs_on_its_own_line(self):
		self.assertEqual(MessageComposer("KSF Kitengela").sign("Hello"), "Hello\nKSF Kitengela")

	def test_no_signature_leaves_message_alone(self):
		self.assertEqual(MessageComposer("").sign("Hello  "), "Hello")

	def test_does_not_sign_twice(self):
		self.assertEqual(MessageComposer("KSF").sign("Thanks\nKSF"), "Thanks\nKSF")

	def test_empty_message_stays_empty(self):
		self.assertEqual(MessageComposer("KSF").sign(""), "")

	def test_compose_fills_variables_then_signs(self):
		self.assertEqual(MessageComposer("KSF").compose("Hi {name}", {"name": "Jane"}), "Hi Jane\nKSF")

	def test_reads_signature_from_settings(self):
		frappe.db.set_single_value("UPEOSMS Settings", "sms_signature", "From Settings")
		self.assertEqual(MessageComposer.from_settings().signature, "From Settings")


class TestSenderProfile(IntegrationTestCase):
	def test_own_keys(self):
		profile = SenderProfile(TEXTSMS_KEYS, TEXTSMS_KEYS)
		self.assertEqual(profile.as_dict(), {"sender_id": "TESTSENDER", "source": "own"})

	def test_inherited_keys_are_shared(self):
		self.assertEqual(SenderProfile(TEXTSMS_KEYS, {}).source, SenderProfile.SHARED)

	def test_no_keys_is_missing(self):
		self.assertEqual(SenderProfile({}, {}).source, SenderProfile.MISSING)


class TestSampleRecipientFile(IntegrationTestCase):
	def test_sample_reads_back_as_valid_upload(self):
		path = frappe.get_site_path("private", "files", "usms_sample_test.xlsx")
		with open(path, "wb") as f:
			f.write(SampleRecipientFile().as_xlsx())
		self.addCleanup(lambda: __import__("os").remove(path))

		rows, columns = _read_xlsx(path)

		self.assertEqual(columns, ["mobile", "name", "amount"])
		self.assertEqual(rows[0]["mobile"], "0712345678")
		self.assertEqual(len(rows), len(SampleRecipientFile.ROWS))


class TestQuickSend(IntegrationTestCase):
	def setUp(self):
		frappe.db.set_single_value("UPEOSMS Settings", "sms_signature", "KSF Test")
		self.before = set(frappe.get_all("SMS Campaign", pluck="name"))

	def tearDown(self):
		for name in set(frappe.get_all("SMS Campaign", pluck="name")) - self.before:
			frappe.db.delete("SMS Send Log", {"campaign": name})
			frappe.db.delete("SMS Recipient", {"campaign": name})
			frappe.db.delete("SMS Campaign", {"name": name})
		frappe.db.set_single_value("UPEOSMS Settings", "sms_signature", None)
		frappe.db.commit()

	def _send(self, numbers, message="Test message"):
		return QuickSend(numbers, message, MessageComposer.from_settings()).send()

	def test_sends_signed_message_to_each_number(self):
		with (
			patch.dict(frappe.local.conf, TEXTSMS_KEYS),
			patch("upeosms.api.sms._post", return_value=TEXTSMS_OK) as post,
		):
			result = self._send("0712345678, 0798765432\n0712345678")

		self.assertEqual([r.status for r in result["results"]], ["Sent", "Sent"])
		self.assertEqual(post.call_count, 2)
		self.assertEqual(post.call_args.args[1]["message"], "Test message\nKSF Test")
		self.assertEqual(frappe.db.count("SMS Send Log", {"campaign": result["campaign"]}), 2)
		self.assertEqual(frappe.db.get_value("SMS Campaign", result["campaign"], "status"), "Completed")

	def test_reports_provider_failure_per_number(self):
		failed = {"responses": [{"response-description": "Invalid sender id", "response-code": 1001}]}
		with patch.dict(frappe.local.conf, TEXTSMS_KEYS), patch("upeosms.api.sms._post", return_value=failed):
			result = self._send("0712345678")

		self.assertEqual(result["results"][0].status, "Failed")
		self.assertIn("Invalid sender id", result["results"][0].error_message)

	def test_rejects_bad_input_before_sending(self):
		cases = [
			("", "Test"),
			("0712345678", ""),
			("12345", "Test"),
			("0712345678", "Hi {name}"),
			(" ".join(f"07123456{i:02d}" for i in range(QuickSend.LIMIT + 1)), "Test"),
		]
		with patch("upeosms.api.sms._post") as post:
			for numbers, message in cases:
				with (
					self.subTest(numbers=numbers, message=message),
					self.assertRaises(frappe.ValidationError),
				):
					self._send(numbers, message)
		post.assert_not_called()


class TestConsoleEndpoints(IntegrationTestCase):
	def test_preview_is_signed(self):
		frappe.db.set_single_value("UPEOSMS Settings", "sms_signature", "KSF Test")
		self.addCleanup(frappe.db.set_single_value, "UPEOSMS Settings", "sms_signature", None)
		preview = page._build_preview([{"mobile": "0712345678", "name": "Jane"}], "Hi {name}")
		self.assertEqual(preview[0]["message"], "Hi Jane\nKSF Test")

	def test_context_shows_sender_and_signature(self):
		with patch.dict(frappe.local.conf, TEXTSMS_KEYS):
			ctx = page.get_console_context()
		self.assertEqual(ctx["sender"]["sender_id"], "TESTSENDER")
		self.assertEqual(ctx["quick_send_limit"], QuickSend.LIMIT)

	def test_endpoints_refuse_users_who_cannot_run_campaigns(self):
		with self.set_user("Guest"), patch("upeosms.api.sms._post") as post:
			for call in (
				page.get_console_context,
				page.download_sample_file,
				lambda: page.quick_send("0712345678", "Hi"),
				lambda: page.get_campaign_progress_from_page("x"),
			):
				with self.assertRaises(frappe.PermissionError):
					call()
		post.assert_not_called()


class TestConsoleSettings(IntegrationTestCase):
	def tearDown(self):
		frappe.db.set_single_value("UPEOSMS Settings", "sms_signature", None)

	def test_saves_tidied_signature(self):
		self.assertEqual(ConsoleSettings().save_signature("  KSF   Kitengela "), "KSF Kitengela")
		self.assertEqual(MessageComposer.from_settings().signature, "KSF Kitengela")

	def test_empty_clears_signature(self):
		ConsoleSettings().save_signature("KSF")
		self.assertEqual(ConsoleSettings().save_signature(""), "")
		self.assertEqual(MessageComposer.from_settings().signature, "")

	def test_rejects_long_signature(self):
		with self.assertRaises(frappe.ValidationError):
			ConsoleSettings().save_signature("x" * (ConsoleSettings.MAX_SIGNATURE + 1))

	def test_needs_settings_write_permission(self):
		with self.set_user("Guest"), self.assertRaises(frappe.PermissionError):
			ConsoleSettings().save_signature("KSF")


class TestCampaignReport(IntegrationTestCase):
	def setUp(self):
		self.campaign = frappe.get_doc(
			{"doctype": "SMS Campaign", "campaign_name": "Report test", "status": "Completed with Errors"}
		).insert(ignore_permissions=True)
		statuses = ["Sent"] * 3 + ["Failed"] + ["Pending"]
		for idx, status in enumerate(statuses, start=1):
			frappe.get_doc(
				{
					"doctype": "SMS Recipient",
					"campaign": self.campaign.name,
					"row_index": idx,
					"mobile": f"07000000{idx:02d}",
					"status": status,
					"error_message": "Traceback (most recent call last):\nValueError: Invalid Kenyan mobile number"
					if status == "Failed"
					else None,
				}
			).insert(ignore_permissions=True)

	def tearDown(self):
		frappe.db.delete("SMS Recipient", {"campaign": self.campaign.name})
		frappe.db.delete("SMS Campaign", {"name": self.campaign.name})

	def test_counts_by_status(self):
		report = CampaignReport(self.campaign.name).as_dict()
		self.assertEqual(report["counts"], {"Sent": 3, "Failed": 1, "Pending": 1})
		self.assertEqual(len(report["recipients"]), 5)

	def test_filters_and_shows_last_error_line(self):
		report = CampaignReport(self.campaign.name).as_dict(status="Failed")
		self.assertEqual([r.status for r in report["recipients"]], ["Failed"])
		self.assertEqual(report["recipients"][0].error_message, "ValueError: Invalid Kenyan mobile number")

	def test_unknown_filter_shows_everyone(self):
		self.assertEqual(len(CampaignReport(self.campaign.name).as_dict(status="Nope")["recipients"]), 5)

	def test_pages(self):
		with patch.object(CampaignReport, "PAGE_SIZE", 2):
			report = CampaignReport(self.campaign.name)
			self.assertEqual(
				[r.mobile for r in report.as_dict(start=2)["recipients"]], ["0700000003", "0700000004"]
			)

	def test_missing_campaign(self):
		with self.assertRaises(frappe.DoesNotExistError):
			CampaignReport("no-such-campaign")

	def test_endpoint_refuses_guests(self):
		with self.set_user("Guest"), self.assertRaises(frappe.PermissionError):
			page.get_campaign_detail(self.campaign.name)

	def test_results_export_lists_every_recipient(self):
		from io import BytesIO

		from openpyxl import load_workbook

		report = CampaignReport(self.campaign.name)
		sheet = load_workbook(BytesIO(report.results_xlsx())).active
		rows = list(sheet.iter_rows(values_only=True))
		self.assertEqual(rows[0][:3], ("Mobile", "Name", "Status"))
		self.assertEqual(len(rows), 6)
		self.assertEqual(rows[4][4], "ValueError: Invalid Kenyan mobile number")
		self.assertEqual(report.results_filename, "report_test_results.xlsx")


class TestCampaignRetry(IntegrationTestCase):
	def setUp(self):
		self.campaign = frappe.get_doc(
			{"doctype": "SMS Campaign", "campaign_name": "Retry test", "status": "Completed with Errors"}
		).insert(ignore_permissions=True)
		for idx, status in enumerate(["Sent", "Failed", "Failed"], start=1):
			frappe.get_doc(
				{
					"doctype": "SMS Recipient",
					"campaign": self.campaign.name,
					"row_index": idx,
					"mobile": "0712345678",
					"status": status,
					"error_message": "boom" if status == "Failed" else None,
				}
			).insert(ignore_permissions=True)

	def tearDown(self):
		frappe.db.delete("SMS Recipient", {"campaign": self.campaign.name})
		frappe.db.delete("SMS Campaign", {"name": self.campaign.name})
		frappe.db.commit()

	def test_requeues_only_failed(self):
		with patch("upeosms.services.campaign_retry.enqueue_campaign_send") as enqueue:
			self.assertEqual(CampaignRetry(self.campaign.name).run(), 2)
		enqueue.assert_called_once_with(self.campaign.name)
		statuses = frappe.get_all("SMS Recipient", {"campaign": self.campaign.name}, pluck="status")
		self.assertEqual(sorted(statuses), ["Pending", "Pending", "Sent"])
		self.assertEqual(frappe.db.get_value("SMS Campaign", self.campaign.name, "status"), "Queued")

	def test_refuses_while_sending(self):
		self.campaign.db_set("status", "Sending")
		with patch("upeosms.services.campaign_retry.enqueue_campaign_send") as enqueue:
			with self.assertRaises(frappe.ValidationError):
				CampaignRetry(self.campaign.name).run()
		enqueue.assert_not_called()

	def test_refuses_when_nothing_failed(self):
		frappe.db.set_value("SMS Recipient", {"campaign": self.campaign.name}, "status", "Sent")
		with self.assertRaises(frappe.ValidationError):
			CampaignRetry(self.campaign.name).run()


class TestCampaignHistory(IntegrationTestCase):
	def setUp(self):
		self.names = [
			frappe.get_doc({"doctype": "SMS Campaign", "campaign_name": name, "status": status})
			.insert(ignore_permissions=True)
			.name
			for name, status in (("History alpha", "Completed"), ("History beta", "Failed"))
		]

	def tearDown(self):
		frappe.db.delete("SMS Campaign", {"name": ["in", self.names]})

	def test_search_and_status(self):
		names = [c.campaign_name for c in CampaignHistory("History").page()["campaigns"]]
		self.assertIn("History alpha", names)
		self.assertIn("History beta", names)
		failed = [c.campaign_name for c in CampaignHistory("History", "Failed").page()["campaigns"]]
		self.assertEqual(failed, ["History beta"])

	def test_unknown_status_is_ignored(self):
		self.assertEqual(len(CampaignHistory("History", "Bogus").page()["campaigns"]), 2)

	def test_has_more(self):
		with patch.object(CampaignHistory, "PAGE_SIZE", 1):
			page_one = CampaignHistory("History").page()
		self.assertEqual(len(page_one["campaigns"]), 1)
		self.assertTrue(page_one["has_more"])
